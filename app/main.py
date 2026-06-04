from __future__ import annotations

import importlib.util
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from fastapi import FastAPI, UploadFile, File, HTTPException, Request, Response, Body
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parent.parent
APP_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
DATA_DIR = BASE_DIR / "data"
STATE_FILE = DATA_DIR / "state.json"
DICTIONARY_FILE = DATA_DIR / "standard_dictionary.json"
DEFAULT_DICTIONARY_FILE = APP_DIR / "pruefturbo_standard_dictionary.json"
for d in (UPLOAD_DIR, OUTPUT_DIR, DATA_DIR):
    d.mkdir(exist_ok=True)

ENGINE_PATH = APP_DIR / "pruefturbo_engine.py"
spec = importlib.util.spec_from_file_location("pruefturbo_engine", ENGINE_PATH)
engine = importlib.util.module_from_spec(spec)
sys.modules["pruefturbo_engine"] = engine
assert spec and spec.loader
spec.loader.exec_module(engine)

app = FastAPI(title="PruefTurbo Web Completed", version="2.0.0")
app.mount("/static", StaticFiles(directory=str(APP_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

# Session-local runtime state. For local single-user use this behaves like the previous app.
# It also prevents different browser sessions from overwriting each other in memory.
SESSIONS: Dict[str, Dict[str, Any]] = {}

class DictionaryEntry(BaseModel):
    category: str = Field(default="General")
    field: str = Field(default="Action")
    entry: str
    translation: str = ""
    notes: str = ""
    source: str = "PruefTurbo_Web_Completed"

class DictionaryUpdate(DictionaryEntry):
    row_id: str

class GeneratePayload(BaseModel):
    selected_keys: List[str] = Field(default_factory=list)

class SaveCaplPayload(BaseModel):
    filename: str = "generated.can"
    content: str

class TestCaseBuilderPayload(BaseModel):
    excel_row: Optional[int] = None
    tc_id: str = ""
    requirement_id: str = ""
    title: str = ""
    author: str = ""
    date: str = ""
    objective: str = ""
    precondition: str = ""
    action: str = ""
    expectation: str = ""
    actual_results: str = ""
    postcondition: str = ""
    test_level: str = ""
    system: str = ""
    priority: str = ""
    ecu: str = ""
    status: str = "Draft"
    reviewer: str = ""
    platform: str = ""
    variant: str = ""
    architecture: str = ""
    constraints: str = ""
    test_platform: str = ""
    test_goal: str = ""
    extra_attrs: Dict[str, str] = Field(default_factory=dict)


def _new_state() -> Dict[str, Any]:
    return {
        "specs_path": None,
        "dictionary_path": str(DICTIONARY_FILE if DICTIONARY_FILE.exists() else DEFAULT_DICTIONARY_FILE),
        "testcases": [],
        "dictionary_rows": _load_default_dictionary_rows(),
        "last_output": None,
        "created_on": datetime.now(timezone.utc).isoformat(),
    }


def _sid(request: Request, response: Optional[Response] = None) -> str:
    sid = request.cookies.get("pt_session_id")
    if not sid or sid not in SESSIONS:
        sid = uuid.uuid4().hex
        SESSIONS[sid] = _new_state()
        if response is not None:
            response.set_cookie("pt_session_id", sid, httponly=True, samesite="lax")
    return sid


def _state(request: Request, response: Optional[Response] = None) -> Dict[str, Any]:
    return SESSIONS[_sid(request, response)]


def _persist_dictionary(rows: List[dict]) -> None:
    grouped: Dict[str, List[dict]] = {}
    for row in rows:
        cat = row.get("category") or "General"
        item = dict(row)
        item.pop("category", None)
        item.pop("row_id", None)
        grouped.setdefault(cat, []).append(item)
    DICTIONARY_FILE.write_text(json.dumps({"schema": "pruefturbo.standard_dictionary.v1", "version": "web-completed", "updated_on": datetime.now(timezone.utc).isoformat(), "dictionary": grouped}, indent=2, ensure_ascii=False), encoding="utf-8")


def _save_upload(upload: UploadFile, prefix: str) -> Path:
    suffix = Path(upload.filename or "uploaded.bin").suffix or ".bin"
    safe_name = f"{prefix}_{uuid.uuid4().hex}{suffix}"
    path = UPLOAD_DIR / safe_name
    with path.open("wb") as f:
        f.write(upload.file.read())
    return path


def _tc_to_dict(tc) -> dict:
    attrs = getattr(tc, "attributes", {}) or {}
    return {
        "row_key": f"row_{tc.excel_row}",
        "excel_row": tc.excel_row,
        "tc_id": getattr(tc, "tc_id", "") or attrs.get("TC ID", ""),
        "title": getattr(tc, "title", "") or attrs.get("TC title", "") or attrs.get("Title", ""),
        "requirements": getattr(tc, "requirements", "") or attrs.get("Requirement ID", ""),
        "priority": getattr(tc, "priority", "") or attrs.get("Priority", ""),
        "ecu": getattr(tc, "ecu", "") or attrs.get("ECU", "") or attrs.get("Device", ""),
        "status": attrs.get("Test Specification Status", "") or attrs.get("Status", ""),
        "valid": bool(getattr(tc, "is_valid_for_generation", True)),
        "block_reason": getattr(tc, "generation_block_reason", ""),
        "capl_name": getattr(tc, "capl_name", ""),
        "attrs": attrs,
    }


def _selected_testcases(st: Dict[str, Any], selected_keys: List[str]):
    testcases = st.get("testcases", [])
    if not selected_keys:
        return []
    selected = set(selected_keys)
    return [tc for tc in testcases if f"row_{tc.excel_row}" in selected or str(tc.excel_row) in selected]


def _flatten_dictionary(dict_by_category) -> List[dict]:
    rows: List[dict] = []
    if isinstance(dict_by_category, dict):
        # Accept wrapped JSON: {schema, version, dictionary:{...}}
        if "dictionary" in dict_by_category and isinstance(dict_by_category["dictionary"], dict):
            dict_by_category = dict_by_category["dictionary"]
        for category, entries in dict_by_category.items():
            if not isinstance(entries, list):
                continue
            for row in entries:
                copied = dict(row)
                copied.setdefault("category", category)
                rows.append(copied)
    elif isinstance(dict_by_category, list):
        rows = [dict(r) for r in dict_by_category]
    return _with_row_ids(rows)


def _with_row_ids(rows: List[dict]) -> List[dict]:
    result = []
    for i, row in enumerate(rows):
        copied = dict(row)
        copied.setdefault("category", copied.get("Category", "General"))
        copied.setdefault("field", copied.get("Field", "Action"))
        copied.setdefault("entry", copied.get("Entry", ""))
        copied.setdefault("translation", copied.get("Translation", ""))
        copied.setdefault("notes", copied.get("Notes", ""))
        base = f"{copied.get('category','')}|{copied.get('field','')}|{copied.get('entry','')}|{i}"
        copied["row_id"] = copied.get("row_id") or uuid.uuid5(uuid.NAMESPACE_URL, base).hex
        result.append(copied)
    return result


def _load_default_dictionary_rows() -> List[dict]:
    path = DICTIONARY_FILE if DICTIONARY_FILE.exists() else DEFAULT_DICTIONARY_FILE
    if not path.exists():
        return []
    try:
        if path.suffix.lower() == ".json":
            return _flatten_dictionary(json.loads(path.read_text(encoding="utf-8")))
    except Exception:
        return []
    return []


def _dictionary_from_upload(path: Path) -> List[dict]:
    if path.suffix.lower() == ".json":
        return _flatten_dictionary(json.loads(path.read_text(encoding="utf-8")))
    # Original function name says Excel, but some desktop versions also support workbook-style dictionaries.
    dict_by_category = engine.read_standard_dictionary_excel(path)
    return _flatten_dictionary(dict_by_category)


def _require_testcases(st: Dict[str, Any]):
    if not st.get("testcases"):
        raise HTTPException(status_code=400, detail="No test specifications loaded yet.")
    return st["testcases"]

def _next_excel_row(st: Dict[str, Any]) -> int:
    rows = [getattr(tc, "excel_row", 1) for tc in st.get("testcases", [])]
    return (max(rows) + 1) if rows else 2


def _payload_to_testcase(payload: TestCaseBuilderPayload, excel_row: Optional[int] = None):
    row_no = excel_row or payload.excel_row or 2
    attrs = dict(payload.extra_attrs or {})
    attrs.update({
        "TC ID": payload.tc_id.strip() or f"WEB_TC_{row_no}",
        "Requirement ID": payload.requirement_id.strip(),
        "TC title": payload.title.strip(),
        "Author": payload.author.strip(),
        "Date": payload.date.strip(),
        "Objective / Scope": payload.objective.strip(),
        "Precondition": payload.precondition.strip(),
        "Action": payload.action.strip(),
        "Expectation": payload.expectation.strip(),
        "Actual Results": payload.actual_results.strip(),
        "Postcondition": payload.postcondition.strip(),
        "Test Level": payload.test_level.strip(),
        "System": payload.system.strip(),
        "Priority": payload.priority.strip(),
        "ECU": payload.ecu.strip(),
        "Test Specification Status": payload.status.strip() or "Draft",
        "Reviewer": payload.reviewer.strip(),
        "Platform": payload.platform.strip(),
        "Variant / Model": payload.variant.strip(),
        "Architecture": payload.architecture.strip(),
        "Test environment constraints": payload.constraints.strip(),
        "Test Platform": payload.test_platform.strip(),
        "Test Goal": payload.test_goal.strip(),
    })
    all_text = "\n".join(str(v) for v in attrs.values())
    valid = not engine.is_placeholder_tc_id(attrs.get("TC ID", ""))
    block_reason = "" if valid else "Placeholder or missing Test Case ID; automatic CAPL generation skipped."
    return engine.TestCaseRow(
        excel_row=row_no,
        attributes=attrs,
        request_bytes=engine.parse_hex_bytes(attrs.get("Action", "")),
        expected_pattern=engine.parse_expected_pattern(attrs.get("Expectation", "")),
        has_timing_check=bool(__import__("re").search(r"timestamp|delta|response\s*time|OBD_GST_CONSTANT_ResponseTime", all_text, flags=__import__("re").IGNORECASE)),
        is_negative=bool(__import__("re").search(r"\b(NRC|negative|7F)\b", all_text, flags=__import__("re").IGNORECASE)),
        is_valid_for_generation=valid,
        generation_block_reason=block_reason,
    )


def _testcase_to_builder_payload(tc, clone: bool = False) -> TestCaseBuilderPayload:
    attrs = getattr(tc, "attributes", {}) or {}
    tc_id = getattr(tc, "tc_id", "") or attrs.get("TC ID", "")
    if clone:
        tc_id = f"{tc_id}_COPY" if tc_id else "WEB_TC_COPY"
    return TestCaseBuilderPayload(
        tc_id=tc_id,
        requirement_id=getattr(tc, "requirements", "") or attrs.get("Requirement ID", ""),
        title=getattr(tc, "title", "") or attrs.get("TC title", ""),
        author=attrs.get("Author", ""),
        date=str(attrs.get("Date", ""))[:10],
        objective=attrs.get("Objective / Scope", "") or attrs.get("Objective", ""),
        precondition=attrs.get("Precondition", "") or attrs.get("Pre Action", ""),
        action=attrs.get("Action", ""),
        expectation=attrs.get("Expectation", "") or attrs.get("Expected Results", ""),
        actual_results=attrs.get("Actual Results", ""),
        postcondition=attrs.get("Postcondition", "") or attrs.get("Post Action", ""),
        test_level=attrs.get("Test Level", ""),
        system=attrs.get("System", ""),
        priority=getattr(tc, "priority", "") or attrs.get("Priority", ""),
        ecu=getattr(tc, "ecu", "") or attrs.get("ECU", ""),
        status=attrs.get("Test Specification Status", "Draft") or attrs.get("Status", "Draft"),
        reviewer=attrs.get("Reviewer", ""),
        platform=attrs.get("Platform", ""),
        variant=attrs.get("Variant / Model", ""),
        architecture=attrs.get("Architecture", ""),
        constraints=attrs.get("Test environment constraints", ""),
        test_platform=attrs.get("Test Platform", ""),
        test_goal=attrs.get("Test Goal", ""),
    )


def _refresh_testcase_names(st: Dict[str, Any]) -> None:
    try:
        engine.make_unique_capl_names(st.get("testcases", []))
    except Exception:
        pass


def _export_testcases_to_excel(testcases, path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSpecs"
    column_map = getattr(engine, "COLUMN_MAP", {})
    max_col = max(column_map.values()) if column_map else 26
    headers = [""] * max_col
    for name, idx in column_map.items():
        headers[idx - 1] = name
    for i, name in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=name or f"Column {i}")
    for out_row, tc in enumerate(testcases, start=2):
        attrs = getattr(tc, "attributes", {}) or {}
        for name, idx in column_map.items():
            ws.cell(row=out_row, column=idx, value=attrs.get(name, ""))
    wb.save(path)
    wb.close()


def _split_requirements(value: str) -> List[str]:
    try:
        return list(engine.split_requirement_ids(value))
    except Exception:
        raw = str(value or "")
        for sep in ["\n", ";", ",", "|"]:
            raw = raw.replace(sep, " ")
        return [x.strip() for x in raw.split() if x.strip()]


def _build_traceability(testcases) -> dict:
    req_to_tc: Dict[str, List[dict]] = {}
    tc_to_req: List[dict] = []
    orphan_testcases: List[dict] = []
    for tc in testcases:
        d = _tc_to_dict(tc)
        reqs = _split_requirements(d.get("requirements", ""))
        tc_to_req.append({"row_key": d["row_key"], "excel_row": d["excel_row"], "tc_id": d["tc_id"], "title": d["title"], "requirements": reqs})
        if not reqs:
            orphan_testcases.append(d)
        for req in reqs:
            req_to_tc.setdefault(req, []).append({"row_key": d["row_key"], "excel_row": d["excel_row"], "tc_id": d["tc_id"], "title": d["title"], "status": d["status"]})
    duplicate_requirements = {req: rows for req, rows in req_to_tc.items() if len(rows) > 1}
    return {
        "requirement_to_testcases": req_to_tc,
        "testcase_to_requirements": tc_to_req,
        "orphan_testcases": orphan_testcases,
        "duplicate_requirement_coverage": duplicate_requirements,
        "total_requirements": len(req_to_tc),
        "total_testcases": len(testcases),
    }


def _calculate_kpi(testcases) -> dict:
    trace = _build_traceability(testcases)
    valid = [_tc_to_dict(tc) for tc in testcases if bool(getattr(tc, "is_valid_for_generation", True))]
    invalid = [_tc_to_dict(tc) for tc in testcases if not bool(getattr(tc, "is_valid_for_generation", True))]
    total = len(testcases)
    covered = trace["total_requirements"]
    orphan = len(trace["orphan_testcases"])
    by_status: Dict[str, int] = {}
    by_priority: Dict[str, int] = {}
    by_ecu: Dict[str, int] = {}
    for tc in testcases:
        d = _tc_to_dict(tc)
        by_status[d.get("status") or "Unspecified"] = by_status.get(d.get("status") or "Unspecified", 0) + 1
        by_priority[d.get("priority") or "Unspecified"] = by_priority.get(d.get("priority") or "Unspecified", 0) + 1
        by_ecu[d.get("ecu") or "Unspecified"] = by_ecu.get(d.get("ecu") or "Unspecified", 0) + 1
    return {
        "total_testcases": total,
        "valid_for_generation": len(valid),
        "blocked_from_generation": len(invalid),
        "total_unique_requirements": covered,
        "orphan_testcases": orphan,
        "requirements_with_multiple_testcases": len(trace["duplicate_requirement_coverage"]),
        "coverage_percent": round(((total - orphan) / total * 100), 2) if total else 0,
        "by_status": by_status,
        "by_priority": by_priority,
        "by_ecu": by_ecu,
        "blocked_rows": invalid,
    }


@app.get("/")
async def index(request: Request, response: Response):
    _sid(request, response)
    return templates.TemplateResponse(request=request, name="index.html", context={"request": request})

@app.get("/api/health")
async def health():
    return {"ok": True, "engine": ENGINE_PATH.name, "version": app.version}

@app.post("/api/session/reset")
async def reset_session(request: Request, response: Response):
    sid = _sid(request, response)
    SESSIONS[sid] = _new_state()
    return {"ok": True}


@app.get("/api/session/save")
async def save_session(request: Request, response: Response):
    st = _state(request, response)
    payload = {
        "schema": "pruefturbo.session.v1",
        "saved_on": datetime.now(timezone.utc).isoformat(),
        "testcases": [_tc_to_dict(tc) for tc in st.get("testcases", [])],
        "dictionary_rows": st.get("dictionary_rows", []),
        "last_output": Path(st.get("last_output", "")).name if st.get("last_output") else None,
    }
    out = OUTPUT_DIR / f"PruefTurbo_session_{uuid.uuid4().hex[:8]}.json"
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return FileResponse(out, filename=out.name, media_type="application/json")

@app.post("/api/upload-specs")
async def upload_specs(request: Request, response: Response, file: UploadFile = File(...)):
    st = _state(request, response)
    path = _save_upload(file, "specs")
    try:
        testcases = engine.load_testcases(path)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to parse test specs: {exc}") from exc
    st["specs_path"] = str(path)
    st["testcases"] = testcases
    return {"ok": True, "filename": file.filename, "count": len(testcases), "testcases": [_tc_to_dict(tc) for tc in testcases]}

@app.get("/api/testcases")
async def get_testcases(request: Request, response: Response):
    st = _state(request, response)
    return {"testcases": [_tc_to_dict(tc) for tc in st.get("testcases", [])]}

@app.post("/api/testcases/builder")
async def builder_add_testcase(request: Request, response: Response, payload: TestCaseBuilderPayload):
    st = _state(request, response)
    tc = _payload_to_testcase(payload, _next_excel_row(st))
    st.setdefault("testcases", []).append(tc)
    _refresh_testcase_names(st)
    return {"ok": True, "testcase": _tc_to_dict(tc), "testcases": [_tc_to_dict(x) for x in st.get("testcases", [])]}

@app.put("/api/testcases/builder/{row_key}")
async def builder_update_testcase(request: Request, response: Response, row_key: str, payload: TestCaseBuilderPayload):
    st = _state(request, response)
    testcases = list(st.get("testcases", []))
    for i, existing in enumerate(testcases):
        existing_key = f"row_{getattr(existing, 'excel_row', '')}"
        if existing_key == row_key or str(getattr(existing, "excel_row", "")) == row_key:
            testcases[i] = _payload_to_testcase(payload, getattr(existing, "excel_row", None))
            st["testcases"] = testcases
            _refresh_testcase_names(st)
            return {"ok": True, "testcase": _tc_to_dict(testcases[i]), "testcases": [_tc_to_dict(x) for x in testcases]}
    raise HTTPException(status_code=404, detail="Test case not found")

@app.post("/api/testcases/builder/{row_key}/clone")
async def builder_clone_testcase(request: Request, response: Response, row_key: str):
    st = _state(request, response)
    testcases = list(st.get("testcases", []))
    for existing in testcases:
        existing_key = f"row_{getattr(existing, 'excel_row', '')}"
        if existing_key == row_key or str(getattr(existing, "excel_row", "")) == row_key:
            payload = _testcase_to_builder_payload(existing, clone=True)
            clone = _payload_to_testcase(payload, _next_excel_row(st))
            st.setdefault("testcases", []).append(clone)
            _refresh_testcase_names(st)
            return {"ok": True, "testcase": _tc_to_dict(clone), "testcases": [_tc_to_dict(x) for x in st.get("testcases", [])]}
    raise HTTPException(status_code=404, detail="Test case not found")

@app.delete("/api/testcases/builder/{row_key}")
async def builder_delete_testcase(request: Request, response: Response, row_key: str):
    st = _state(request, response)
    before = list(st.get("testcases", []))
    after = [tc for tc in before if f"row_{getattr(tc, 'excel_row', '')}" != row_key and str(getattr(tc, "excel_row", "")) != row_key]
    if len(after) == len(before):
        raise HTTPException(status_code=404, detail="Test case not found")
    st["testcases"] = after
    _refresh_testcase_names(st)
    return {"ok": True, "count": len(after), "testcases": [_tc_to_dict(x) for x in after]}

@app.post("/api/testcases/builder/clear")
async def builder_clear_testcases(request: Request, response: Response):
    st = _state(request, response)
    st["testcases"] = []
    return {"ok": True, "count": 0, "testcases": []}

@app.get("/api/testcases/export")
async def export_testcases(request: Request, response: Response):
    st = _state(request, response)
    testcases = _require_testcases(st)
    out = OUTPUT_DIR / f"PruefTurbo_testcases_{uuid.uuid4().hex[:8]}.xlsx"
    _export_testcases_to_excel(testcases, out)
    return FileResponse(out, filename=out.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.post("/api/upload-dictionary")
async def upload_dictionary(request: Request, response: Response, file: UploadFile = File(...)):
    st = _state(request, response)
    path = _save_upload(file, "dictionary")
    try:
        rows = _dictionary_from_upload(path)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to parse dictionary: {exc}") from exc
    st["dictionary_path"] = str(path)
    st["dictionary_rows"] = rows
    _persist_dictionary(rows)
    cats = sorted({r.get("category", "") for r in rows if r.get("category")})
    return {"ok": True, "filename": file.filename, "count": len(rows), "categories": cats, "rows": rows[:500]}

@app.get("/api/dictionary")
async def get_dictionary(request: Request, response: Response):
    st = _state(request, response)
    rows = st.get("dictionary_rows") or _load_default_dictionary_rows()
    st["dictionary_rows"] = rows
    return {"count": len(rows), "categories": sorted({r.get("category", "") for r in rows if r.get("category")}), "rows": rows[:5000]}

@app.post("/api/dictionary/create")
async def create_dictionary(request: Request, response: Response, payload: dict = Body(default={})):
    st = _state(request, response)
    rows = _flatten_dictionary(payload.get("dictionary", payload)) if payload else []
    st["dictionary_rows"] = rows
    _persist_dictionary(rows)
    return {"ok": True, "count": len(rows), "rows": rows}

@app.post("/api/dictionary/add")
async def add_dictionary_entry(request: Request, response: Response, entry: DictionaryEntry):
    st = _state(request, response)
    rows = list(st.get("dictionary_rows") or [])
    new_row = entry.model_dump()
    new_row["row_id"] = uuid.uuid4().hex
    rows.append(new_row)
    st["dictionary_rows"] = rows
    _persist_dictionary(rows)
    return {"ok": True, "row": new_row, "count": len(rows)}

@app.put("/api/dictionary/update/{row_id}")
async def update_dictionary_entry(request: Request, response: Response, row_id: str, payload: dict = Body(...)):
    st = _state(request, response)
    rows = list(st.get("dictionary_rows") or [])
    for i, row in enumerate(rows):
        if row.get("row_id") == row_id:
            updated = dict(row)
            updated.update({k: v for k, v in payload.items() if k != "row_id"})
            updated["row_id"] = row_id
            rows[i] = updated
            st["dictionary_rows"] = rows
            _persist_dictionary(rows)
            return {"ok": True, "row": updated}
    raise HTTPException(status_code=404, detail="Dictionary row not found")

@app.delete("/api/dictionary/delete/{row_id}")
async def delete_dictionary_entry(request: Request, response: Response, row_id: str):
    st = _state(request, response)
    rows = list(st.get("dictionary_rows") or [])
    new_rows = [r for r in rows if r.get("row_id") != row_id]
    if len(new_rows) == len(rows):
        raise HTTPException(status_code=404, detail="Dictionary row not found")
    st["dictionary_rows"] = new_rows
    _persist_dictionary(new_rows)
    return {"ok": True, "count": len(new_rows)}

@app.get("/api/dictionary/export")
async def export_dictionary(request: Request, response: Response):
    st = _state(request, response)
    rows = st.get("dictionary_rows") or []
    _persist_dictionary(rows)
    return FileResponse(DICTIONARY_FILE, filename="pruefturbo_standard_dictionary_web_completed.json", media_type="application/json")

@app.post("/api/generate-direct")
async def generate_direct(request: Request, response: Response, payload: GeneratePayload):
    st = _state(request, response)
    selected = _selected_testcases(st, payload.selected_keys)
    if not selected:
        raise HTTPException(status_code=400, detail="No selected test cases. Select at least one row.")
    try:
        capl = engine.generate_capl(selected)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Direct CAPL generation failed: {exc}") from exc
    out = OUTPUT_DIR / f"PruefTurbo_direct_{uuid.uuid4().hex[:8]}.can"
    out.write_text(capl, encoding="utf-8")
    st["last_output"] = str(out)
    return {"ok": True, "filename": out.name, "download_url": f"/api/download/{out.name}", "selected_count": len(selected), "preview": capl[:10000]}

@app.post("/api/generate-dictionary")
async def generate_dictionary(request: Request, response: Response, payload: GeneratePayload):
    st = _state(request, response)
    selected = _selected_testcases(st, payload.selected_keys)
    if not selected:
        raise HTTPException(status_code=400, detail="No selected test cases. Select at least one row.")
    rows = st.get("dictionary_rows", [])
    if not rows:
        raise HTTPException(status_code=400, detail="No dictionary loaded. Upload or create a dictionary first.")
    try:
        capl = engine.generate_capl_from_dictionary(selected, rows)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Dictionary CAPL generation failed: {exc}") from exc
    out = OUTPUT_DIR / f"PruefTurbo_dictionary_{uuid.uuid4().hex[:8]}.can"
    out.write_text(capl, encoding="utf-8")
    st["last_output"] = str(out)
    return {"ok": True, "filename": out.name, "download_url": f"/api/download/{out.name}", "selected_count": len(selected), "preview": capl[:10000]}

@app.post("/api/capl/save")
async def save_capl(payload: SaveCaplPayload):
    safe = Path(payload.filename).name or "generated.can"
    if not safe.lower().endswith((".can", ".cin", ".txt")):
        safe += ".can"
    out = OUTPUT_DIR / safe
    out.write_text(payload.content, encoding="utf-8")
    return {"ok": True, "filename": out.name, "download_url": f"/api/download/{out.name}"}

@app.get("/api/download/{filename}")
async def download(filename: str):
    path = OUTPUT_DIR / Path(filename).name
    if not path.exists():
        raise HTTPException(status_code=404, detail="Output file not found")
    return FileResponse(path, filename=path.name, media_type="application/octet-stream")

@app.get("/api/traceability")
async def traceability(request: Request, response: Response):
    st = _state(request, response)
    return _build_traceability(_require_testcases(st))

@app.get("/api/traceability/export")
async def export_traceability(request: Request, response: Response, mode: str = "downstream"):
    st = _state(request, response)
    trace = _build_traceability(_require_testcases(st))
    wb = openpyxl.Workbook()
    ws = wb.active
    direction = "upstream" if str(mode).lower().startswith("up") else "downstream"
    ws.title = "Upstream" if direction == "upstream" else "Downstream"
    if direction == "upstream":
        ws.append(["Test Case ID", "Title", "Excel Row", "Requirement ID"] )
        for row in trace.get("testcase_to_requirements", []):
            reqs = row.get("requirements") or [""]
            for req in reqs:
                ws.append([row.get("tc_id", ""), row.get("title", ""), row.get("excel_row", ""), req])
    else:
        ws.append(["Requirement ID", "Test Case ID", "Title", "Excel Row", "Status"] )
        for req, tcs in (trace.get("requirement_to_testcases") or {}).items():
            for tc in tcs:
                ws.append([req, tc.get("tc_id", ""), tc.get("title", ""), tc.get("excel_row", ""), tc.get("status", "")])
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 14), 60)
    out = OUTPUT_DIR / f"PruefTurbo_traceability_{direction}_{uuid.uuid4().hex[:8]}.xlsx"
    wb.save(out)
    wb.close()
    return FileResponse(out, filename=out.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/api/kpi")
async def kpi(request: Request, response: Response):
    st = _state(request, response)
    return _calculate_kpi(_require_testcases(st))
